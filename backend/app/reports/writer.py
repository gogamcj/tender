"""XLSX·CSV 보고서 작성.

모든 보고서에는 '보고서 정보' 영역(생성일시, 분석 기준일, 데이터 범위, 산식, 적용 규칙, 추정값 여부, 면책 문구)을 넣는다.
- 금액은 정수로, 비율은 소수로 쓰되 원래 문자열 값을 바꾸지 않는다(Decimal → 숫자 변환은 표시용).
- 수식 주입 방지: =, +, -, @, 탭, 줄바꿈으로 시작하는 글자 값은 앞에 작은따옴표를 붙인다(숫자 값은 제외).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.importers.parsing import KST

DANGEROUS = ("=", "+", "-", "@", "\t", "\r", "\n")


@dataclass
class Table:
    title: str
    columns: list[tuple[str, str]]  # (키, 머리글)
    rows: list[dict[str, Any]]
    kinds: dict[str, str] = field(default_factory=dict)  # 키 → money / rate / int / text


@dataclass
class Report:
    name: str
    info: list[tuple[str, Any]]  # 보고서 정보(항목, 값)
    tables: list[Table]
    generated_at: datetime = field(default_factory=lambda: datetime.now(UTC))

    def info_rows(self) -> list[tuple[str, Any]]:
        return [("보고서", self.name), ("생성일시", self.generated_at.astimezone(KST).strftime("%Y-%m-%d %H:%M:%S KST")), *self.info]


def safe_text(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in DANGEROUS:
        try:
            Decimal(value)  # 음수 숫자 문자열은 그대로 둔다
            return value
        except (InvalidOperation, ValueError):
            return "'" + value
    return value


def _cell(value: Any, kind: str | None) -> Any:
    if value is None or value == "":
        return None
    if kind in {"money", "int"}:
        try:
            d = Decimal(str(value))
            return int(d) if d == d.to_integral_value() else float(d)
        except (InvalidOperation, ValueError):
            return safe_text(str(value))
    if kind == "rate":
        try:
            return float(Decimal(str(value)))
        except (InvalidOperation, ValueError):
            return safe_text(str(value))
    if isinstance(value, bool):
        return "예" if value else "아니오"
    if isinstance(value, list | tuple):
        return safe_text("; ".join(str(v) for v in value))
    if isinstance(value, dict):
        return safe_text("; ".join(f"{k}={v}" for k, v in value.items()))
    return safe_text(str(value)) if not isinstance(value, int | float) else value


HEADER_FILL = PatternFill("solid", fgColor="DDE4EE")


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "보고서 정보"
    for k, v in report.info_rows():
        ws.append([k, _cell(v, None)])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    used: set[str] = {ws.title}
    for t in report.tables:
        title = t.title[:31]
        i = 2
        while title in used:
            title = f"{t.title[:28]}_{i}"
            i += 1
        used.add(title)
        sh = wb.create_sheet(title)
        sh.append([h for _, h in t.columns])
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = HEADER_FILL
        for r in t.rows:
            sh.append([_cell(r.get(k), t.kinds.get(k)) for k, _ in t.columns])
        for idx, (k, h) in enumerate(t.columns, start=1):
            letter = get_column_letter(idx)
            sh.column_dimensions[letter].width = max(10, min(50, len(h) * 2 + 4))
            kind = t.kinds.get(k)
            if kind in {"money", "int"}:
                for cell in sh[letter][1:]:
                    cell.number_format = "#,##0"
            elif kind == "rate":
                for cell in sh[letter][1:]:
                    cell.number_format = "0.000###"
        sh.freeze_panes = "A2"
        if not t.rows:
            sh.append(["(자료 없음)"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(report: Report, table_index: int = 0) -> bytes:
    """CSV 는 표 하나만 담는다. 보고서 정보는 앞쪽 '#' 줄로 넣는다(Excel 호환 UTF-8 BOM)."""
    buf = io.StringIO()
    w = csv.writer(buf)
    for k, v in report.info_rows():
        w.writerow([f"# {k}", _cell(v, None)])
    w.writerow([])
    t = report.tables[table_index]
    w.writerow([h for _, h in t.columns])
    for r in t.rows:
        w.writerow(["" if (c := _cell(r.get(k), t.kinds.get(k))) is None else c for k, _ in t.columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")
