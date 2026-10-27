"""업로드 파일 읽기: XLSX / XLS / CSV.

모든 셀은 문자열로 읽어 원본 그대로 보존한다(형식 변환은 parsing 단계에서 명시적으로 한다).
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
CSV_ENCODINGS = ("utf-8-sig", "cp949")  # cp949 는 EUC-KR 의 상위 집합
ALLOWED_EXTENSIONS = {"xlsx", "xls", "csv"}


class FileRejected(ValueError):
    """업로드 파일 검사 실패."""


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[str]]  # 헤더 제외, 셀은 문자열(빈 셀은 "")

    @property
    def row_count(self) -> int:
        return len(self.rows)


@dataclass
class ParsedFile:
    file_name: str
    file_hash: str
    source_type: str
    encoding: str | None
    sheets: list[Sheet] = field(default_factory=list)

    def sheet(self, name: str | None) -> Sheet:
        if name is None:
            return self.sheets[0]
        for s in self.sheets:
            if s.name == name:
                return s
        raise FileRejected(f"시트를 찾을 수 없습니다: {name}")


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _extension(file_name: str) -> str:
    return file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ") if value.time() != time(0) else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        # 엑셀 숫자는 float 로 온다. 정수로 표현 가능한 값은 정수 문자열로, 아니면 repr 로(가장 짧은 정확 표현).
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int | Decimal):
        return str(value)
    return str(value).strip()


def _header(raw: list[str]) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(raw):
        name = (h or "").strip() or f"열{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        names.append(name)
    return names


def _build_sheet(name: str, matrix: list[list[str]], header_row: int) -> Sheet:
    if len(matrix) <= header_row:
        return Sheet(name=name, columns=[], rows=[])
    data_rows = [r for r in matrix[header_row + 1 :] if any(c.strip() for c in r)]  # 완전히 빈 행 제외
    if len(data_rows) > MAX_ROWS:
        raise FileRejected(f"행이 너무 많습니다(최대 {MAX_ROWS:,}행).")
    raw_header = list(matrix[header_row])
    # 헤더보다 긴 데이터 행이 있으면 잘라내지 않고 '열N' 이름의 열을 추가한다.
    width = max([len(raw_header)] + [len(r) for r in data_rows])

    def column_empty(k: int) -> bool:
        in_header = k < len(raw_header) and raw_header[k].strip()
        return not in_header and all(k >= len(r) or not r[k].strip() for r in data_rows)

    while width > 0 and column_empty(width - 1):
        width -= 1  # 헤더와 값이 모두 빈 꼬리 열 제거(엑셀 서식만 남은 열)
    header = _header((raw_header + [""] * width)[:width])
    rows = [(r + [""] * width)[:width] for r in data_rows]
    return Sheet(name=name, columns=header, rows=rows)


def _read_csv(data: bytes, header_row: int) -> tuple[Sheet, str]:
    if b"\x00" in data[:8192]:
        raise FileRejected("CSV 파일에 이진 데이터가 포함되어 있습니다.")
    for enc in CSV_ENCODINGS:
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileRejected("CSV 인코딩을 확인할 수 없습니다(UTF-8, CP949/EUC-KR 만 지원).")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    matrix = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    return _build_sheet("CSV", matrix, header_row), enc


def _read_xlsx(data: bytes, header_row: int) -> list[Sheet]:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise FileRejected("올바른 XLSX 파일이 아닙니다.")
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = set(zf.namelist())
        if "[Content_Types].xml" not in names:
            raise FileRejected("올바른 XLSX 파일이 아닙니다.")
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            raise FileRejected("매크로가 포함된 파일은 업로드할 수 없습니다.")
        # 압축 폭탄 방지: 압축 해제 크기 합계 제한
        if sum(i.file_size for i in zf.infolist()) > MAX_FILE_BYTES * 20:
            raise FileRejected("압축 해제 크기가 너무 큽니다.")
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            matrix = []
            for row in ws.iter_rows(values_only=True):
                matrix.append([_cell_to_str(v) for v in row])
                if len(matrix) > MAX_ROWS + header_row + 1:
                    raise FileRejected(f"행이 너무 많습니다(최대 {MAX_ROWS:,}행).")
            sheets.append(_build_sheet(ws.title, matrix, header_row))
    finally:
        wb.close()
    return sheets


def _read_xls(data: bytes, header_row: int) -> list[Sheet]:
    if not data.startswith(b"\xd0\xcf\x11\xe0"):
        raise FileRejected("올바른 XLS 파일이 아닙니다.")
    try:
        import xlrd  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - 선택 의존성
        raise FileRejected("XLS 지원 패키지(xlrd)가 설치되어 있지 않습니다. XLSX 또는 CSV 로 저장해 올려 주세요.") from exc
    book = xlrd.open_workbook(file_contents=data)
    sheets = []
    for sh in book.sheets():
        matrix = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell_to_str(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                else:
                    row.append(_cell_to_str(cell.value))
            matrix.append(row)
        sheets.append(_build_sheet(sh.name, matrix, header_row))
    return sheets


def read_upload(file_name: str, data: bytes, *, header_row: int = 0) -> ParsedFile:
    if not data:
        raise FileRejected("빈 파일입니다.")
    if len(data) > MAX_FILE_BYTES:
        raise FileRejected(f"파일이 너무 큽니다(최대 {MAX_FILE_BYTES // 1024 // 1024}MB).")
    ext = _extension(file_name)
    if ext not in ALLOWED_EXTENSIONS:
        raise FileRejected(f"지원하지 않는 형식입니다: .{ext} (XLSX, XLS, CSV 만 가능)")
    parsed = ParsedFile(file_name=file_name, file_hash=file_hash(data), source_type=ext, encoding=None)
    if ext == "csv":
        sheet, enc = _read_csv(data, header_row)
        parsed.sheets = [sheet]
        parsed.encoding = enc
    elif ext == "xlsx":
        parsed.sheets = _read_xlsx(data, header_row)
    else:
        parsed.sheets = _read_xls(data, header_row)
    if not parsed.sheets:
        raise FileRejected("읽을 수 있는 시트가 없습니다.")
    return parsed
