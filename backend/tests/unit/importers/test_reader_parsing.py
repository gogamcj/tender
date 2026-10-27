import io
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest

from app.engines.calc import RateUnit
from app.importers import reader
from app.importers.fields import suggest_mapping
from app.importers.parsing import ParseError, parse_date, parse_datetime, parse_int, parse_money, parse_rate
from app.importers.reader import FileRejected, read_upload

CSV_TEXT = "공고번호,공고명,기초금액\nA-1,가상 공사,\"1,000,000\"\nA-2,가상 용역,2000000\n"


# ------------------------------------------------------------------ CSV 인코딩


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp949", "euc-kr"])
def test_csv_encodings(encoding):
    parsed = read_upload("a.csv", CSV_TEXT.encode(encoding))
    sheet = parsed.sheets[0]
    assert sheet.columns == ["공고번호", "공고명", "기초금액"]
    assert sheet.rows[0] == ["A-1", "가상 공사", "1,000,000"]
    assert parsed.encoding in ("utf-8-sig", "cp949")


def test_csv_tab_delimited():
    parsed = read_upload("a.csv", CSV_TEXT.replace(",", "\t").replace('"1\t000\t000"', "1000000").encode())
    assert parsed.sheets[0].columns == ["공고번호", "공고명", "기초금액"]


def test_csv_unknown_encoding_rejected():
    with pytest.raises(FileRejected):
        read_upload("a.csv", "공고".encode("utf-16"))


def test_csv_binary_rejected():
    with pytest.raises(FileRejected):
        read_upload("a.csv", b"abc\x00def")


def test_empty_and_unsupported_files():
    with pytest.raises(FileRejected):
        read_upload("a.csv", b"")
    with pytest.raises(FileRejected):
        read_upload("a.exe", b"MZ")
    with pytest.raises(FileRejected):
        read_upload("a.xlsx", b"not a zip")
    with pytest.raises(FileRejected):
        read_upload("a.xls", b"not ole")


def test_file_size_limit(monkeypatch):
    monkeypatch.setattr(reader, "MAX_FILE_BYTES", 10)
    with pytest.raises(FileRejected):
        read_upload("a.csv", b"x" * 11)


def test_header_handling_duplicates_blank_and_extra_columns():
    text = "공고번호,,공고번호,\nA,1,2,,\nB,3,4,5,extra\n\n"
    sheet = read_upload("a.csv", text.encode()).sheets[0]
    assert sheet.columns == ["공고번호", "열2", "공고번호_2", "열4", "열5"]
    assert sheet.rows == [["A", "1", "2", "", ""], ["B", "3", "4", "5", "extra"]]


def test_header_row_option():
    text = "제목: 가상 자료\n공고번호,공고명\nA,B\n"
    sheet = read_upload("a.csv", text.encode(), header_row=1).sheets[0]
    assert sheet.columns == ["공고번호", "공고명"]


def test_xlsx_reading_keeps_values():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "결과"
    ws.append(["공고번호", "공고일", "개찰일시", "기초금액", "낙찰률", None])
    ws.append(["A-1", date(2024, 3, 4), datetime(2024, 3, 20, 11, 0), 150000000, 87.745, None])
    wb.create_sheet("빈시트")
    buf = io.BytesIO()
    wb.save(buf)
    parsed = read_upload("a.xlsx", buf.getvalue())
    assert [s.name for s in parsed.sheets] == ["결과", "빈시트"]
    s = parsed.sheet("결과")
    assert s.columns == ["공고번호", "공고일", "개찰일시", "기초금액", "낙찰률"]  # 꼬리 빈 열 제거
    assert s.rows[0] == ["A-1", "2024-03-04", "2024-03-20 11:00:00", "150000000", "87.745"]
    with pytest.raises(FileRejected):
        parsed.sheet("없는시트")


def test_xls_reading():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("S")
    for c, v in enumerate(["공고번호", "기초금액"]):
        ws.write(0, c, v)
    ws.write(1, 0, "A-1")
    ws.write(1, 1, 1000000)
    buf = io.BytesIO()
    wb.save(buf)
    s = read_upload("a.xls", buf.getvalue()).sheets[0]
    assert s.rows == [["A-1", "1000000"]]


def test_macro_xlsx_rejected():
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/vbaProject.bin", "x")
    with pytest.raises(FileRejected):
        read_upload("a.xlsx", buf.getvalue())


def test_suggest_mapping_exact_names_only():
    m = suggest_mapping(["입찰공고번호", "공고 명", "기초금액(원)", "기초금액비고", "낙찰률"])
    assert m["notice_number"] == "입찰공고번호"
    assert m["title"] == "공고 명"
    assert m["base_amount"] == "기초금액(원)"
    assert m["winning_rate"] == "낙찰률"
    assert "기초금액비고" not in m.values()


# ------------------------------------------------------------------ 값 변환


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("1,234,567", "1234567"), ("1234567원", "1234567"), ("₩ 1,000", "1000"), ("0", "0"), ("1000.0", "1000")],
)
def test_parse_money_ok(raw, expected):
    assert parse_money(raw) == Decimal(expected)


@pytest.mark.parametrize("raw", ["1.2억", "3천만", "abc", "-100", "100.5", "1,2,3a"])
def test_parse_money_rejects(raw):
    with pytest.raises(ParseError):
        parse_money(raw)


def test_parse_rate_units_and_warnings():
    assert parse_rate("87.745", RateUnit.PERCENT) == (Decimal("87.745000"), [])
    assert parse_rate("0.87745", RateUnit.RATIO)[0] == Decimal("87.745000")
    assert parse_rate("87.745%", RateUnit.RATIO)[0] == Decimal("87.745000")  # % 표기는 단위 설정보다 우선
    _, w = parse_rate("0.87745", RateUnit.PERCENT)
    assert w  # 소수처럼 보이면 경고
    _, w2 = parse_rate("87.745", RateUnit.RATIO)
    assert w2
    with pytest.raises(ParseError):
        parse_rate("abc", RateUnit.PERCENT)
    with pytest.raises(ParseError):
        parse_rate("87.1234567", RateUnit.PERCENT)


@pytest.mark.parametrize(
    "raw", ["2024-03-04", "2024.03.04", "2024. 3. 4.", "2024/3/4", "20240304", "2024년 3월 4일", "2024-03-04 00:00:00"]
)
def test_parse_date_formats(raw):
    assert parse_date(raw) == date(2024, 3, 4)


@pytest.mark.parametrize("raw", ["2024-02-30", "2023-13-01", "04-03-2024", "어제", "2024-3"])
def test_parse_date_invalid(raw):
    with pytest.raises(ParseError):
        parse_date(raw)


def test_parse_datetime_kst_to_utc():
    assert parse_datetime("2024-03-20 11:00") == datetime(2024, 3, 20, 2, 0, tzinfo=UTC)
    assert parse_datetime("202403201100") == datetime(2024, 3, 20, 2, 0, tzinfo=UTC)
    assert parse_datetime("2024-03-20") == datetime(2024, 3, 19, 15, 0, tzinfo=UTC)
    with pytest.raises(ParseError):
        parse_datetime("2024-03-20 25:00")


def test_parse_int():
    assert parse_int("1,234") == 1234
    assert parse_int("12개") == 12
    with pytest.raises(ParseError):
        parse_int("1.5")


def test_suggest_mapping_skips_other_units():
    m = suggest_mapping(["공고번호", "기초금액(천원)", "낙찰률(%)"])
    assert "base_amount" not in m
    assert m["winning_rate"] == "낙찰률(%)"
