import io
from datetime import UTC, datetime, timedelta
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy.orm import sessionmaker

from app.db.models import BidResult, ReservePrice

from .test_analysis_api import AS_OF, history  # noqa: F401


def _notice(client, number, **kw):
    r = client.post("/api/notices", json={"notice_number": number, "title": f"공고 {number}", **kw})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_dashboard(make_client):
    c = make_client("user", "dash")
    soon = (datetime.now(UTC) + timedelta(days=2)).isoformat()
    later = (datetime.now(UTC) + timedelta(days=30)).isoformat()
    a = _notice(c, "D-1", bid_deadline=soon)
    _notice(c, "D-2", bid_deadline=later)
    c.post("/api/favorites", json={"notice_id": a})
    d = c.get("/api/dashboard").json()
    assert d["notices"]["total"] == 2
    assert [n["id"] for n in d["upcoming_deadlines"]] == [a] and d["upcoming_deadlines"][0]["favorite"] is True
    assert d["favorites"] == 1 and d["api_key_configured"] in (True, False)
    assert {"total", "unverified"} <= set(d["rules"])


def test_favorites_are_private(make_client):
    a, b = make_client("user", "fa"), make_client("user", "fb")
    nid = _notice(a, "F-1")
    fav = a.post("/api/favorites", json={"notice_id": nid, "memo": "검토"}).json()
    assert a.post("/api/favorites", json={"notice_id": nid}).status_code == 409
    assert b.get("/api/favorites").json() == []
    assert b.patch(f"/api/favorites/{fav['id']}", json={"memo": "x"}).status_code == 404
    assert a.patch(f"/api/favorites/{fav['id']}", json={"memo": "수정"}).json()["memo"] == "수정"
    assert a.delete(f"/api/favorites/{fav['id']}").status_code == 204
    assert a.get("/api/favorites").json() == []


def test_saved_filters(make_client):
    c = make_client("user", "sf")
    r = c.post("/api/saved-filters", json={"name": "우리 지역", "criteria": {"region": "가상도"}})
    assert r.status_code == 201
    assert c.post("/api/saved-filters", json={"name": "우리 지역", "criteria": {}}).status_code == 409
    assert c.post("/api/saved-filters", json={"name": "x", "filter_type": "bad", "criteria": {}}).status_code == 400
    assert [s["name"] for s in c.get("/api/saved-filters").json()] == ["우리 지역"]


def test_reserve_average_calc(client):
    r = client.post("/api/calc/reserve-average", json={"selected_prices": ["100100000", "99900000", "100300000", "99800000"],
                                                       "base_amount": "100000000"}).json()
    assert r["value"] == "100025000" and r["price_status"] == "scenario"
    assert r["sajeong"]["value"].startswith("0.025")
    assert client.post("/api/calc/reserve-average", json={"selected_prices": []}).status_code == 422


def test_notice_reserve_prices_and_stats(client, engine):
    nid = _notice(client, "RP-1", procurement_type="construction", issuing_authority_type="local", contract_method="제한경쟁",
                  award_method="적격심사", base_amount="100000000", opening_date="2026-01-10T11:00:00+09:00")
    factory = sessionmaker(bind=engine)
    with factory() as s:
        res = BidResult(notice_id=nid, source_name="g2b_scsbid", rebid_number=0, actual_base_price=Decimal("100000000"),
                        actual_estimated_price=Decimal("100025000"), actual_sajeong_rate=Decimal("0.025"),
                        opening_datetime=datetime(2026, 1, 10, 2, tzinfo=UTC))
        s.add(res)
        s.flush()
        amounts = ["100100000", "99900000", "100300000", "99800000"] + [str(98000000 + i * 250000) for i in range(11)]
        for i, a in enumerate(amounts, start=1):
            s.add(ReservePrice(result_id=res.id, seq=i, amount=Decimal(a), drawn=i <= 4, draw_count=None, source_name="g2b_scsbid"))
        s.commit()
    d = client.get(f"/api/notices/{nid}/reserve-prices").json()
    r0 = d["results"][0]
    assert d["available"] and len(r0["prices"]) == 15 and r0["drawn_count"] == 4
    assert r0["prices"][0]["rate_vs_base"] == "0.100000"
    assert r0["estimated_check"]["consistent"] is True
    st = client.post("/api/analysis/reserve-prices", json={"filters": {"period": None, "as_of": AS_OF}}).json()
    assert len(st["strata"]) == 1
    s0 = st["strata"][0]
    assert s0["notice_count"] == 1 and s0["all_prices"]["n"] == 15 and s0["drawn_prices"]["n"] == 4
    assert [x["seq"] for x in s0["draws_by_seq"]] == [1, 2, 3, 4]
    assert "예측하지 않습니다" in st["disclaimer"]
    empty = client.get(f"/api/notices/{_notice(client, 'RP-2')}/reserve-prices").json()
    assert empty["available"] is False and empty["note"]


def _xlsx(resp):
    assert resp.status_code == 200, resp.text
    return load_workbook(io.BytesIO(resp.content))


def test_export_notices_xlsx_and_csv(history):  # noqa: F811
    wb = _xlsx(history.get("/api/reports/notices.xlsx", params={"agency": "샘플시청", "has_result": True}))
    info = {r[0].value: r[1].value for r in wb["보고서 정보"].iter_rows()}
    assert info["보고서"] == "공고 목록" and "생성일시" in info and "agency" in info["검색 조건"]
    sh = wb["공고 목록"]
    header = [c.value for c in sh[1]]
    assert header[0] == "공고번호" and "사정률(%)" in header
    base_col = header.index("기초금액") + 1
    assert isinstance(sh.cell(row=2, column=base_col).value, int)
    csv = history.get("/api/reports/notices.csv", params={"agency": "샘플시청"})
    text = csv.content.decode("utf-8")
    assert text.startswith("﻿# 보고서") and "공고번호" in text
    assert "attachment" in csv.headers["content-disposition"]
    assert history.get("/api/reports/notices.pdf").status_code == 400


def test_export_runs(history):  # noqa: F811
    run = history.post("/api/analysis/stats", json={"filters": {"period": None, "as_of": AS_OF}, "metric": "sajeong_rate",
                                                    "group_by": ["agency_name"], "save": True}).json()["analysis_run_id"]
    wb = _xlsx(history.get(f"/api/reports/runs/{run}.xlsx"))
    assert wb.sheetnames[:3] == ["보고서 정보", "조건별 요약", "묶음별"]
    info = {r[0].value: r[1].value for r in wb["보고서 정보"].iter_rows()}
    assert info["데이터 버전"].startswith("n=") and info["분석 기준일·기간"]
    nid = history.post("/api/notices", json={"notice_number": "EXP-1", "title": "대상", "procurement_type": "construction",
                                             "issuing_authority_type": "local", "contract_method": "제한경쟁", "award_method": "적격심사",
                                             "base_amount": "800000000", "bid_deadline": "2026-10-30T10:00:00+09:00"}).json()["id"]
    rec = history.post("/api/recommend", json={"notice_id": nid, "period": None, "lower_limit_rate_pct": "87.745", "save": True}).json()
    wb = _xlsx(history.get(f"/api/reports/runs/{rec['analysis_run_id']}.xlsx"))
    sh = wb["시나리오 비교"]
    assert [c.value for c in sh[1]][:2] == ["시나리오", "방법"] and sh.max_row == 4
    info = {r[0].value: r[1].value for r in wb["보고서 정보"].iter_rows()}
    assert "가정 기반" in info["추정값 여부"] and info["하한 산식"]
    assert history.get("/api/reports/runs/99999.xlsx").status_code == 404


def test_export_cost_access_and_formula_injection(make_client):
    a, b = make_client("user", "ca"), make_client("user", "cb")
    cid = a.post("/api/costs", json={"material_cost": "100000000", "labor_cost": "0", "equipment_cost": "0", "subcontract_cost": "0",
                                     "transport_cost": "0", "site_overhead": "0", "general_overhead": "0", "statutory_costs": "0",
                                     "other_costs": "0", "vat_treatment": "excluded", "notes": "=HYPERLINK(\"x\")"}).json()["id"]
    wb = _xlsx(a.get(f"/api/reports/costs/{cid}.xlsx", params={"bid_amounts": ["110000000", "121000000"], "bid_vat_included": True}))
    rows = list(wb["투찰금액별 손익"].iter_rows(values_only=True))
    assert rows[1][3] == 0 and rows[2][3] == 10000000
    assert b.get(f"/api/reports/costs/{cid}.xlsx").status_code == 404
    nid = _notice(a, "=CMD|1")
    text = a.get("/api/reports/notices.csv", params={"q": "CMD"}).content.decode("utf-8")
    assert "'=CMD|1" in text and str(nid)


def test_export_rules(client):
    client.post("/api/rules", json={"rule_code": "R1", "rule_name": "규칙", "effective_from": "2020-01-01",
                                    "formula_definition": {"module": "simple_lower_limit", "params": {"lower_limit_rate_pct": "87.745"}}})
    wb = _xlsx(client.get("/api/reports/rules.xlsx"))
    assert wb["규칙"].cell(row=2, column=1).value == "R1"
